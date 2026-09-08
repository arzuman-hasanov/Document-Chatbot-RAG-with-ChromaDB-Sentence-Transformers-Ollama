"""Concrete parsers for supported document formats."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from pypdf import PdfReader
from docx import Document as DocxDocument

from .base import DocumentParser
from .exceptions import DocumentParsingError, EmptyDocumentError
from .models import Document


def _normalize_metadata(value: Any) -> Any:
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    if isinstance(value, list):
        return [_normalize_metadata(item) for item in value]
    if isinstance(value, dict):
        return {str(key): _normalize_metadata(item) for key, item in value.items()}
    return str(value)


class TxtParser(DocumentParser):
    file_type = "txt"

    def parse(self, file_path: Path) -> Document:
        if not file_path.exists():
            raise DocumentParsingError(f"Text file not found: {file_path}")

        raw_bytes = file_path.read_bytes()
        if not raw_bytes:
            raise EmptyDocumentError(f"Text document is empty: {file_path.name}")

        for encoding in ("utf-8-sig", "utf-16", "latin-1"):
            try:
                text = raw_bytes.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise DocumentParsingError(f"Unable to decode text file: {file_path.name}")

        cleaned_text = text.strip()
        if not cleaned_text:
            raise EmptyDocumentError(f"Text document contains no usable text: {file_path.name}")

        metadata = {
            "size_bytes": len(raw_bytes),
            "line_count": len(cleaned_text.splitlines()),
            "encoding": encoding,
        }
        return Document(
            source_filename=file_path.name,
            file_type=self.file_type,
            extracted_text=cleaned_text,
            metadata=metadata,
        )


class PdfParser(DocumentParser):
    file_type = "pdf"

    def parse(self, file_path: Path) -> Document:
        if not file_path.exists():
            raise DocumentParsingError(f"PDF file not found: {file_path}")

        try:
            reader = PdfReader(str(file_path))
        except Exception as exc:  # pragma: no cover - broad guard for malformed PDFs
            raise DocumentParsingError(f"Unable to open PDF file: {file_path.name}") from exc

        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:  # pragma: no cover - decryption errors are reported clearly
                raise DocumentParsingError(f"PDF is encrypted and cannot be read: {file_path.name}") from exc

        pages: list[str] = []
        for page in reader.pages:
            content = page.extract_text() or ""
            if content.strip():
                pages.append(content.strip())

        extracted_text = "\n\n".join(pages)
        if not extracted_text:
            raise EmptyDocumentError(f"PDF document contains no extractable text: {file_path.name}")

        metadata = {
            "pages": len(reader.pages),
            "is_encrypted": reader.is_encrypted,
            "pdf_metadata": _normalize_metadata(dict(reader.metadata or {})),
        }

        return Document(
            source_filename=file_path.name,
            file_type=self.file_type,
            extracted_text=extracted_text,
            metadata=metadata,
        )


class DocxParser(DocumentParser):
    file_type = "docx"

    def parse(self, file_path: Path) -> Document:
        if not file_path.exists():
            raise DocumentParsingError(f"DOCX file not found: {file_path}")

        try:
            document = DocxDocument(str(file_path))
        except Exception as exc:  # pragma: no cover - invalid files are surfaced as clear errors
            raise DocumentParsingError(f"Unable to open DOCX file: {file_path.name}") from exc

        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text and paragraph.text.strip()]
        extracted_text = "\n".join(paragraphs)
        if not extracted_text:
            raise EmptyDocumentError(f"DOCX document contains no text: {file_path.name}")

        metadata = {
            "paragraph_count": len(paragraphs),
            "section_count": len(document.sections),
            "core_properties": _normalize_metadata(
                {
                    "author": document.core_properties.author,
                    "title": document.core_properties.title,
                    "subject": document.core_properties.subject,
                    "keywords": document.core_properties.keywords,
                    "created": str(document.core_properties.created) if document.core_properties.created else None,
                    "modified": str(document.core_properties.modified) if document.core_properties.modified else None,
                }
            ),
        }

        return Document(
            source_filename=file_path.name,
            file_type=self.file_type,
            extracted_text=extracted_text,
            metadata=metadata,
        )
