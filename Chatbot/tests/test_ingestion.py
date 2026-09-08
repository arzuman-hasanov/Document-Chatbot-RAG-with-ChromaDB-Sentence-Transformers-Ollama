from __future__ import annotations

from pathlib import Path

import pytest
from pypdf import PdfWriter
from docx import Document as DocxDocument

from document_chatbot.ingestion import DocumentIngestionService
from document_chatbot.ingestion.exceptions import EmptyDocumentError, UnsupportedDocumentFormatError


@pytest.fixture
def service() -> DocumentIngestionService:
    return DocumentIngestionService()


def _escape_pdf_text(value: str) -> str:
    return value.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _build_minimal_pdf(text: str) -> bytes:
    text_stream = f"BT /F1 12 Tf 72 720 Td ({_escape_pdf_text(text)}) Tj ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 300] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(text_stream.encode('latin-1'))} >>\nstream\n{text_stream}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    pdf = "%PDF-1.4\n"
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(pdf.encode("latin-1")))
        pdf += f"{index} 0 obj\n{obj}\nendobj\n"

    xref_start = len(pdf.encode("latin-1"))
    pdf += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n"
    for offset in offsets[1:]:
        pdf += f"{offset:010d} 00000 n \n"
    pdf += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_start}\n%%EOF"
    return pdf.encode("latin-1")


def test_parse_txt_file(service: DocumentIngestionService, tmp_path: Path) -> None:
    file_path = tmp_path / "sample.txt"
    file_path.write_text("This is a sample text document.\nSecond line.", encoding="utf-8")

    document = service.parse_document(file_path)

    assert document.source_filename == "sample.txt"
    assert document.file_type == "txt"
    assert "sample text document" in document.extracted_text
    assert document.metadata["line_count"] >= 2


def test_parse_document_can_use_explicit_source_filename(service: DocumentIngestionService, tmp_path: Path) -> None:
    file_path = tmp_path / "tmp-upload.txt"
    file_path.write_text("Original document content.", encoding="utf-8")

    document = service.parse_document(file_path, source_filename="original-name.txt")

    assert document.source_filename == "original-name.txt"


def test_parse_pdf_file(service: DocumentIngestionService, tmp_path: Path) -> None:
    file_path = tmp_path / "sample.pdf"
    file_path.write_bytes(_build_minimal_pdf("Hello from PDF"))

    document = service.parse_document(file_path)

    assert document.source_filename == "sample.pdf"
    assert document.file_type == "pdf"
    assert "Hello from PDF" in document.extracted_text
    assert document.metadata["pages"] == 1


def test_parse_docx_file(service: DocumentIngestionService, tmp_path: Path) -> None:
    file_path = tmp_path / "sample.docx"
    doc = DocxDocument()
    doc.add_paragraph("Hello from DOCX paragraph 1")
    doc.add_paragraph("Second paragraph")
    doc.save(file_path)

    document = service.parse_document(file_path)

    assert document.source_filename == "sample.docx"
    assert document.file_type == "docx"
    assert "Hello from DOCX paragraph 1" in document.extracted_text
    assert document.metadata["paragraph_count"] == 2


def test_parse_empty_document_raises(service: DocumentIngestionService, tmp_path: Path) -> None:
    empty_file = tmp_path / "empty.txt"
    empty_file.write_text("", encoding="utf-8")

    with pytest.raises(EmptyDocumentError, match="contains no usable text|is empty"):
        service.parse_document(empty_file)


def test_parse_unsupported_format_raises(service: DocumentIngestionService, tmp_path: Path) -> None:
    unsupported = tmp_path / "notes.md"
    unsupported.write_text("This is markdown", encoding="utf-8")

    with pytest.raises(UnsupportedDocumentFormatError, match="Unsupported document format"):
        service.parse_document(unsupported)
